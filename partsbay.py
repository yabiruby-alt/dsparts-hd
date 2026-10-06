#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
동성모터스 PARTS - BMW My DMS 부품부서 대시보드 생성기

실행하면 (기본 = 상시 실행 모드):
  1. 브라우저가 최소화된 채로 한 번 뜨고 My DMS에 접속합니다. 화면엔 안 보입니다.
  2. 로그인이 안 되어 있으면 화면에 알림창이 뜹니다 - 작업표시줄에서 그 창을 열어 로그인(2FA 포함)해주세요.
  3. 로그인되면 그 브라우저를 계속 켜둔 채로 10분마다 자동으로 데이터를 뽑아 대시보드를 갱신하고
     깃허브에 push합니다. 프로그램을 끄지 않는 한 계속 이 상태로 돌아갑니다 (Ctrl+C로 종료).

`python partsbay.py --once` 로 실행하면 한 번만 갱신하고 바로 종료합니다 (테스트용).

세션(쿠키)은 authdata 폴더에 저장되어, 프로그램을 재시작해도 세션이 유지되는 동안은 재로그인이 필요 없습니다.
동시에 두 개를 띄우면 DMS가 세션을 끊어버리므로 절대 두 인스턴스를 같이 실행하지 마세요(잠금 파일로 방지됨).
"""
import csv
import io
import json
import mimetypes
import os
import queue
import re
import subprocess
import sys
import threading
import time
import urllib.parse
import urllib.request
import webbrowser

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass
from collections import defaultdict
from datetime import datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from jinja2 import Environment, FileSystemLoader
from playwright.sync_api import sync_playwright, Frame, Page

# ============================================================
# 설정 - 지점마다 이 블록만 바꿔서 배포
# ============================================================
BASE_URL = "https://bmwdms.co.kr"
DEALER_CD = "001672"
BIZ_AREA_CD = "051"
BRCH_CD = "15"
BRANCH_NAME = "AS_부산(해운중동)"

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
DOCS_DIR = BASE_DIR / "docs"  # GitHub Pages가 서빙하는 폴더
# 팀 캘린더: 이 이름의 이미지 파일을 docs 폴더에 넣어두면 "팀 캘린더" 메뉴에 그대로 표시됨.
CALENDAR_IMAGE_CANDIDATES = ["team-calendar.jpg", "team-calendar.jpeg", "team-calendar.png", "team-calendar.webp"]
GIT_REMOTE_URL = "https://github.com/yabiruby-alt/dsparts-hd.git"
PUBLISH_TO_GITHUB = True
# 크로미움이 한글 등 비-ASCII 경로의 --user-data-dir 에서 불안정하게 종료되는 문제가 있어
# 로그인 세션 저장 폴더는 항상 ASCII 전용 경로(사용자 홈 폴더 아래)에 둔다.
AUTH_DIR = Path.home() / "AppData" / "Local" / "PartsBayDMS" / "authdata"
TEMPLATE_NAME = "template.html.j2"

# 진행RO 사유: 이 PC에서만 입력(로컬 서버) -> reasons.json 저장 -> 사이트에 반영해 폰/다른 기기는 조회 전용
REASONS_FILE = BASE_DIR / "reasons.json"
REASON_PORT = 8765
ALLOWED_ORIGINS = {"https://yabiruby-alt.github.io", "http://127.0.0.1:8765", "http://localhost:8765"}
REASONS_LOCK = threading.Lock()   # reasons.json 읽기/쓰기
# 타이어 사이즈 조회 마스터: 이 엑셀 파일을 저장하면 매 주기 자동으로 다시 읽어 반영 (사이트에서 직접 수정하지 않음)
# docs 폴더 안에 둬서 원본 파일(타이어 마스터 폴더)이 없어져도 안전하게 보관 — 깃허브에는 올라가지 않도록 .gitignore 처리됨
TIRE_XLSX_PATH = BASE_DIR / "docs" / "타이어_마스터파일.xlsx"
TIRE_FILE = BASE_DIR / "tire_master.json"   # 엑셀을 못 읽을 때 쓸 마지막 성공본 캐시
TIRE_LOCK = threading.Lock()
# 디데이 현황: 구글 폼(제목/목표일/메모) 응답 시트를 CSV로 내보낸 링크를 읽어와 매 주기 자동 반영
# 설정 전까지는 빈 문자열로 두면 "연동 전" 상태로 표시됨
DDAY_SHEET_CSV_URL = "https://docs.google.com/spreadsheets/d/1OmZxvQjz1DutkjB1Jkm7JRPinr_7P60BXBLycYo0RU8/export?format=csv&gid=782790237"
DDAY_FILE = BASE_DIR / "dday_items.json"   # 시트를 못 읽을 때 쓸 마지막 성공본 캐시
DDAY_LOCK = threading.Lock()
# 매출 대시보드 모듈 위치 (지점장 전용 — 파츠베이 공개 사이트와 별도 저장소로 배포)
SALES_DIR = BASE_DIR.parent / "업무자동화 생성" / "매출데이터 사이트"
SALES_ENABLED = (SALES_DIR / "sales.py").exists()
RR_ENABLED = True   # 롤스로이스 DMS 연동 — 켜면 데몬 시작 시 RR 로그인 창이 하나 더 뜸
RR_BASE_URL = "https://www.rrdms.co.kr"   # 롤스로이스 DMS (같은 My DMS 플랫폼, 별도 계정/세션)
RR_AUTH_DIR = Path.home() / "AppData" / "Local" / "PartsBayDMS" / "authdata_rr"
STOCK_QUEUE = queue.Queue()       # 부품번호 재고조회 요청(HTTP 스레드 -> 메인 스레드의 DMS 세션으로 처리)
PUBLISH_LOCK = threading.Lock()   # 렌더 + docs 쓰기 + 깃허브 push 직렬화
LAST_DATA = {}                    # 마지막 갱신 데이터(사유만 바뀌었을 때 DMS 재조회 없이 재렌더용)

WEEKDAYS_KO = ["월", "화", "수", "목", "금", "토", "일"]

# ============================================================
# 사용자 알림 (창은 항상 최소화되어 있으므로, 사람 개입이 필요할 때만
# 화면에 알림창을 띄운다)
# ============================================================

def notify_user(title: str, message: str) -> None:
    safe_title = title.replace("'", "''")
    safe_msg = message.replace("'", "''")
    try:
        subprocess.Popen(
            ["powershell", "-NoProfile", "-WindowStyle", "Hidden", "-Command",
             f"Add-Type -AssemblyName System.Windows.Forms; "
             f"[System.Windows.Forms.MessageBox]::Show('{safe_msg}', '{safe_title}', "
             f"[System.Windows.Forms.MessageBoxButtons]::OK, [System.Windows.Forms.MessageBoxIcon]::Warning) | Out-Null"],
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception as e:
        print(f"[알림 실패] {e}")


# ============================================================
# 브라우저 / 로그인
# ============================================================

def ensure_logged_in(page: Page, timeout_sec: int = 420) -> None:
    page.goto(f"{BASE_URL}/selectHome.do")
    deadline = time.time() + timeout_sec
    notified = False
    last_debug = 0.0
    while time.time() < deadline:
        try:
            # 좌측 메뉴(#icon-parts 등)는 iframe이 아니라 최상위 페이지 자체에 있음. 접혀 있어도
            # DOM에는 존재하므로 visible 상태를 요구하지 않고 attached 여부만 확인한다.
            page.wait_for_selector("#icon-parts", timeout=2500, state="attached")
            print(f"로그인 감지됨. 현재 URL: {page.url}")
            return
        except Exception:
            pass
        if not notified:
            print("DMS 로그인이 필요합니다. 브라우저 창에서 로그인(2FA 포함)해주세요...")
            notify_user(
                "동성모터스 PARTS — 로그인 필요",
                "DMS 세션이 만료되어 대시보드 자동 갱신이 멈춰 있습니다.\n"
                "작업표시줄에서 최소화된 브라우저 창을 열어 로그인(2FA 포함)해주세요.",
            )
            notified = True
        if time.time() - last_debug > 10:
            last_debug = time.time()
            try:
                print(f"[대기중] url={page.url}")
                debug_path = OUTPUT_DIR / "_debug_login.png"
                OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
                page.screenshot(path=str(debug_path))
            except Exception as e:
                print(f"[디버그 스크린샷 실패] {e}")
        page.wait_for_timeout(1500)
    raise TimeoutError("로그인 대기 시간을 초과했습니다. 프로그램을 다시 실행해주세요.")


def click_menu(page: Page, section_id: str, label: str) -> None:
    page.evaluate(
        """({sectionId, label}) => {
            const a = [...document.querySelectorAll('#' + sectionId + ' a[data-url]')]
                .find(x => x.textContent.trim() === label);
            if (!a) throw new Error('메뉴를 찾지 못했습니다: ' + label);
            a.click();
        }""",
        {"sectionId": section_id, "label": label},
    )


def wait_for_frame(page: Page, url_substr: str, timeout_ms: int = 20000) -> Frame:
    deadline = time.time() + timeout_ms / 1000
    while time.time() < deadline:
        for f in page.frames:
            if url_substr in f.url:
                return f
        page.wait_for_timeout(300)
    raise TimeoutError(f"화면을 찾지 못했습니다: {url_substr}")


FETCH_JS = """
async ({url, body}) => {
  return await new Promise((resolve) => {
    jQuery.ajax({
      url, type: 'POST', dataType: 'json', contentType: 'application/json',
      data: JSON.stringify(body),
      success: (res) => resolve({ok: true, rows: res.data || [], total: res.total}),
      error: (xhr) => resolve({ok: false, status: xhr.status, text: (xhr.responseText || '').slice(0, 300)})
    });
  });
}
"""


FETCH_RAW_JS = """
async ({url, body}) => {
  return await new Promise((resolve) => {
    jQuery.ajax({
      url, type: 'POST', dataType: 'text', contentType: 'application/json',
      data: JSON.stringify(body),
      success: (res) => resolve({ok: true, text: res}),
      error: (xhr) => resolve({ok: false, status: xhr.status, text: (xhr.responseText || '').slice(0, 300)})
    });
  });
}
"""


def lookup_stock(page: Page, pn: str) -> dict:
    """부품번호 하나의 해운대(우리 지점) 재고 + RDC 재고 원본 응답. DMS 세션을 쓰므로 메인 스레드에서만 호출."""
    out = {"pn": pn}
    click_menu(page, "icon-parts", "현재고리스트 조회")
    frame = wait_for_frame(page, "selectInventListMain")
    own_body = {
        "recordCountPerPage": 100, "pageIndex": 1, "firstIndex": 0, "lastIndex": 100,
        "sCorpCd": DEALER_CD, "sBizAreaCd": BIZ_AREA_CD, "sBrchCd": BRCH_CD,
        "sProdType": "", "sItemCd": pn, "sItemNm": "", "sStrgeCd": "", "sCrtQtyYn": False,
    }
    try:
        out["own"] = fetch_rows(frame, "/parts/inventory/selectInventoryList.do", own_body)
    except Exception as e:
        out["own_error"] = str(e)[:300]
    rdc_body = {
        "recordCountPerPage": 30, "pageIndex": 1, "firstIndex": 0, "lastIndex": 30,
        "sItemCds": [pn], "sRefDocNo": [], "sRefDocLineNo": [], "sView": "saleView",
    }
    res = frame.evaluate(FETCH_RAW_JS, {"url": "/parts/cmm/popup/selectRdcPartsPopup.do", "body": rdc_body})
    if res.get("ok"):
        try:
            out["rdc"] = json.loads(res["text"])
        except ValueError:
            out["rdc_text"] = (res["text"] or "")[:3000]
    else:
        out["rdc_error"] = res
    return out


def serve_stock_requests(page: Page) -> None:
    while True:
        try:
            pn, holder, done = STOCK_QUEUE.get_nowait()
        except queue.Empty:
            return
        try:
            holder["result"] = lookup_stock(page, pn)
        except Exception as e:
            holder["error"] = f"{type(e).__name__}: {e}"[:300]
        done.set()


def fetch_rows(frame: Frame, url: str, body: dict) -> list:
    result = frame.evaluate(FETCH_JS, {"url": url, "body": body})
    if not result.get("ok"):
        raise RuntimeError(f"요청 실패 {url}: {result}")
    return result["rows"]


# ============================================================
# 데이터 추출
# ============================================================

def extract_receiving(page: Page, today_str: str) -> list:
    click_menu(page, "icon-parts", "입고현황")
    frame = wait_for_frame(page, "selectWhStatusMain")
    body = {
        "recordCountPerPage": 5000, "pageIndex": 1,
        "sRealWhDtFrom": today_str, "sRealWhDtTo": today_str,
        "sBpNm": "", "sItemCd": "", "sWhNo": "", "sPurcOrderNo": "",
        "sAloisCd": "", "sWhTp": "", "sWhStat": "", "sGrnNoFrom": "", "sGrnNoTo": "",
    }
    return fetch_rows(frame, "/parts/whmng/selectWhStatusMain.do", body)


def extract_receiving_range(page: Page, start_str: str, end_str: str) -> list:
    """입고현황을 날짜 범위(예: 이번달 1일~오늘)로 조회. O파트 입출고 내역 메뉴용."""
    click_menu(page, "icon-parts", "입고현황")
    frame = wait_for_frame(page, "selectWhStatusMain")
    body = {
        "recordCountPerPage": 200000, "pageIndex": 1,
        "sRealWhDtFrom": start_str, "sRealWhDtTo": end_str,
        "sBpNm": "", "sItemCd": "", "sWhNo": "", "sPurcOrderNo": "",
        "sAloisCd": "", "sWhTp": "", "sWhStat": "", "sGrnNoFrom": "", "sGrnNoTo": "",
    }
    return fetch_rows(frame, "/parts/whmng/selectWhStatusMain.do", body)


def extract_inventory(page: Page) -> list:
    click_menu(page, "icon-parts", "현재고리스트 조회")
    frame = wait_for_frame(page, "selectInventListMain")
    body = {
        "recordCountPerPage": 200000, "pageIndex": 1, "firstIndex": 0, "lastIndex": 200000,
        "sCorpCd": DEALER_CD, "sBizAreaCd": BIZ_AREA_CD, "sBrchCd": BRCH_CD,
        "sProdType": "", "sItemCd": "", "sItemNm": "", "sStrgeCd": "", "sCrtQtyYn": True,
    }
    return fetch_rows(frame, "/parts/inventory/selectInventoryList.do", body)


def _turnover_body(start_str: str, end_str: str) -> dict:
    return {
        "recordCountPerPage": 200000, "pageIndex": 1,
        "sSearchStartDt": f"{start_str}T00:00:00.000Z",
        "sSearchEndDt": f"{end_str}T23:59:59.000Z",
        "sBrands": [], "sSeriesList": [], "sCarNo": "", "sVinNo": "",
        "sCustTp": "", "sCustNo": "", "sCustNm": "",
        "sDlrCd": DEALER_CD, "sBrchCdList": [BRCH_CD], "sSaList": [], "sRoDocNo": "",
        "sCalcTpCds": [], "sInvcNum": "", "sInvStatCd": "", "sCalcNo": "",  # 빈 배열 = 전체 정산유형(C/I/W/S/CP/WS 등)
        "sIctTradeYn": "", "sItemTpCd": "", "sItemCd": "", "sItemNm": "", "sProdType": "전체",
        "sAloiscd": "", "sRclCampnCdYn": "", "sCampnYn": "", "sCupnCdYn": "", "sSvcTypess": [],
    }


def extract_turnover(page: Page, month_start: str, today_str: str) -> list:
    click_menu(page, "icon-report", "Turn Over 리포트")
    frame = wait_for_frame(page, "selectDLRTurnOver")
    return fetch_rows(frame, "/rpt/raw/selectDLRTurnOver.do", _turnover_body(month_start, today_str))


def extract_resv_sbs(page: Page, stat: str, start_str: str) -> list:
    """서비스예약현황에서 예약상태(01 예약접수/03 예약취소/04 No Show)별 SB. start_str부터 무한대(2099)까지."""
    click_menu(page, "icon-service", "서비스예약현황")
    frame = wait_for_frame(page, "selectResvAcptStatusMain")
    body = {
        "recordCountPerPage": 200000, "pageIndex": 1, "firstIndex": 0, "lastIndex": 200000,
        "sDlrCd": DEALER_CD, "sBizAreaCd": BIZ_AREA_CD, "sBrchCd": BRCH_CD,
        "sResvStartDt": start_str, "sResvEndDt": "2099-12-31", "sResvNo": "", "sChrgSaNm": "",
        "sSvcType": "", "sAcptDstin": "", "sCarRegNo": "", "sVinNo": "", "sCustNm": "", "sCustNo": "",
        "sOnCheckInYn": "", "sResvStat": stat,
    }
    return fetch_rows(frame, "/ser/resvAcpt/selectResvAcptStatus.do", body)


def extract_carin_sbs_without_ro(page: Page, req_rows: list, today) -> tuple:
    """차량접수 상태 SB 중 미처리 부품요청이 걸려있고 RO가 아직 없는 SB. (해당 SB 행 리스트, 최근 180일 차량접수 SB 총수)"""
    start = (today - timedelta(days=180)).strftime("%Y-%m-%d")
    sbs = extract_resv_sbs(page, "02", start)
    cand_ids = {r.get("refDocNo") for r in req_rows if r.get("refDocTp") == "SB"}
    cand = [r for r in sbs if r.get("resvNo") in cand_ids]
    if not cand:
        return [], len(sbs)

    ro_from = min(((r.get("carAcptDtime") or r.get("resvDtime") or "")[:10] or start) for r in cand)
    cur = datetime.strptime(ro_from, "%Y-%m-%d").date() - timedelta(days=2)
    click_menu(page, "icon-report", "RO 리포트")
    frame = wait_for_frame(page, "selectRawRptByRoRptMain")
    resv_with_ro = set()
    while cur <= today:  # RO 리포트는 길게 조회하면 느려서 월 단위로 쪼갬
        nxt = (cur.replace(day=28) + timedelta(days=4)).replace(day=1)
        end = min(nxt - timedelta(days=1), today)
        body = {
            "recordCountPerPage": 200000, "pageIndex": 1, "firstIndex": 0, "lastIndex": 200000,
            "sSelectDt": "csltEndDt", "sSearchStartDt": cur.strftime("%Y-%m-%d"), "sSearchEndDt": end.strftime("%Y-%m-%d"),
            "sBrands": [], "sVinNo": "", "sCarNo": "", "sSeriesList": [], "sCustTp": "", "sCustNo": "", "sCustNm": "",
            "sDlrCd": DEALER_CD, "sBrchCdList": [BRCH_CD], "sSaList": [], "sTechNms": [], "sRoNo": "",
            "sRoStat": "", "sSvcTypes": [], "sCustCarYn": "", "sDqRcrNoti": "", "sCalcTpCd": "", "sCalcTpCds": [],
            "sRclTcApply": "", "sCampnApply": "", "sCupnApply": "", "sLeadTimeUnit": "D", "sExcludeCaseTargetYn": "N",
            "totCnt": 0,
        }
        for r in fetch_rows(frame, "/rpt/raw/selectRawRoRptList.do", body):
            if r.get("resvNo"):
                resv_with_ro.add(r["resvNo"])
        cur = nxt
    return [r for r in cand if r.get("resvNo") not in resv_with_ro], len(sbs)


OPEN_RO_STAT_CODES = ("01", "02", "03", "08", "04", "05")  # 06=인보이스 완료, 07=RO취소 제외


def extract_open_ros(page: Page, today_str: str) -> list:
    """RO 리포트에서 2000-01-01~오늘 발행된 RO 중 인보이스 안 된 것(상태별로 나눠 조회)."""
    click_menu(page, "icon-report", "RO 리포트")
    frame = wait_for_frame(page, "selectRawRptByRoRptMain")
    rows = []
    for stat in OPEN_RO_STAT_CODES:
        body = {
            "recordCountPerPage": 200000, "pageIndex": 1, "firstIndex": 0, "lastIndex": 200000,
            "sSelectDt": "csltEndDt", "sSearchStartDt": "2000-01-01", "sSearchEndDt": today_str,
            "sBrands": [], "sVinNo": "", "sCarNo": "", "sSeriesList": [], "sCustTp": "", "sCustNo": "", "sCustNm": "",
            "sDlrCd": DEALER_CD, "sBrchCdList": [BRCH_CD], "sSaList": [], "sTechNms": [], "sRoNo": "",
            "sRoStat": stat, "sSvcTypes": [], "sCustCarYn": "", "sDqRcrNoti": "", "sCalcTpCd": "", "sCalcTpCds": [],
            "sRclTcApply": "", "sCampnApply": "", "sCupnApply": "", "sLeadTimeUnit": "D", "sExcludeCaseTargetYn": "N",
            "totCnt": 0,
        }
        rows += fetch_rows(frame, "/rpt/raw/selectRawRoRptList.do", body)
    return rows


def extract_part_requests(page: Page) -> list:
    """출고요청관리: 미처리 부품 출고요청(어떤 RO/SB/SP에 재고가 묶여있는지 참조문서번호 포함).
    sNotProcQty(처리잔여수량 존재여부) 필터를 걸면 partStatCd=R("가용(입고)") 건이 빠지므로 필터를 풀어서(빈 값) 전체를 가져온다."""
    click_menu(page, "icon-parts", "출고요청관리")
    frame = wait_for_frame(page, "selectDlvReqMngMain")
    rows = []
    for req_tp in ("01", "03"):  # 01=RO/SB/SP 대부분, 03=소수 잔여 유형(확인됨)
        body = {
            "recordCountPerPage": 5000, "pageIndex": 1, "firstIndex": 0, "lastIndex": 5000,
            "sRefDocNo": "", "sReqStartDt": "2020-01-01", "sReqEndDt": "2099-12-31",  # SB 등 미래 예약분도 놓치지 않게 상한을 사실상 무제한으로
            "sStatCd": "01", "sReqDocNo": "", "sReqUsrId": "", "sPartNo": "",
            "sPartStatCd": "", "sReqBrchCd": "", "sNotProcQty": "", "sPurcTp": "", "sReqTp": req_tp,
        }
        rows += fetch_rows(frame, "/parts/dlv/dlvReq/selectPartReqInfo.do", body)
    return rows


# ============================================================
# 집계 (JS로 검증한 로직을 그대로 Python으로 옮김)
# ============================================================

GROUP_ORDER = ["A", "L", "O", "I", "S"]
GROUP_LABELS = {"A": "상시재고", "L": "로컬조달", "O": "특수/단종계열", "I": "비이동성", "S": "특수발주"}
PROD_LABELS = {"3": "Car Accessories", "5": "Rims and Complete Wheels", "7": "Accessories (Lifestyle)"}


def num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def group_alois(rows, val_fn, qty_fn=None):
    by_code = defaultdict(lambda: {"val": 0.0, "qty": 0.0})
    for r in rows:
        code = r.get("aloisCd") or "?"
        by_code[code]["val"] += val_fn(r)
        if qty_fn:
            by_code[code]["qty"] += qty_fn(r)
    total = sum(v["val"] for v in by_code.values()) or 1.0
    groups = []
    for g in GROUP_ORDER:
        codes = sorted(c for c in by_code if c and c[0] == g)
        if not codes:
            continue
        grows, sub_val, sub_qty = [], 0.0, 0.0
        for c in codes:
            v, q = by_code[c]["val"], by_code[c]["qty"]
            sub_val += v
            sub_qty += q
            row = {"code": c, "val": round(v), "pct": round(100 * v / total, 2)}
            if qty_fn:
                row["qty"] = round(q)
            grows.append(row)
        groups.append({
            "code": g, "label": GROUP_LABELS[g], "oseries": g == "O",
            "rows": grows, "subtotal_val": round(sub_val), "subtotal_qty": round(sub_qty),
            "subtotal_pct": round(100 * sub_val / total, 2),
        })
    return groups


def build_recv(rows, today_str):
    rows = [r for r in rows if r.get("whTp") == "일반입고"]  # 반품입고 제외 (금액 상쇄 방지)

    def val_fn(r):
        return num(r.get("purcAmt"))

    total = round(sum(val_fn(r) for r in rows))
    qty_total = round(sum(num(r.get("whQty")) for r in rows))
    groups = group_alois(rows, val_fn)

    by_grn = defaultdict(float)
    for r in rows:
        by_grn[r.get("grnNo") or "(없음)"] += val_fn(r)
    max_val = max(by_grn.values()) if by_grn else 0
    grn_rows = [
        {"grn": g, "val": round(v), "is_top": v == max_val and v > 0}
        for g, v in sorted(by_grn.items(), key=lambda kv: kv[0], reverse=True)
    ]

    top5 = sorted(rows, key=val_fn, reverse=True)[:5]
    top5_out = [{
        "item": r.get("itemCd"), "name": r.get("itemNm"), "alois": r.get("aloisCd"),
        "grn": r.get("grnNo"), "qty": int(num(r.get("whQty"))), "val": round(val_fn(r)),
        "oseries": (r.get("aloisCd") or "")[:1] == "O",
    } for r in top5]

    return {
        "total": total, "date": today_str, "grn_count": len(by_grn),
        "lines": len(rows), "qty": qty_total,
        "alois_groups": groups, "grn_rows": grn_rows, "top5": top5_out,
    }


def build_inventory(rows):
    pw = [r for r in rows if r.get("strgNm") == "부품창고"]

    def val_fn(r):
        return num(r.get("crtQty")) * num(r.get("movPrc"))

    def qty_fn(r):
        return num(r.get("crtQty"))

    total = round(sum(val_fn(r) for r in pw))
    groups = group_alois(pw, val_fn, qty_fn)
    codes = {r.get("aloisCd") for r in pw if r.get("aloisCd")}

    parts_by_code = defaultdict(list)
    for r in pw:
        if r.get("aloisCd") and num(r.get("crtQty")) > 0:
            parts_by_code[r["aloisCd"]].append({
                "item": r.get("itemCd"), "name": r.get("itemNm"),
                "qty": int(num(r.get("crtQty"))), "val": round(val_fn(r)),
            })
    for g in groups:
        for row in g["rows"]:
            row["parts"] = sorted(parts_by_code.get(row["code"], []), key=lambda p: -p["val"])

    inv = {"total": total, "item_count": len(pw), "code_count": len(codes), "alois_groups": groups}
    return inv, pw


def build_oaov(pw_rows, inv_total):
    def val_fn(r):
        return num(r.get("crtQty")) * num(r.get("movPrc"))

    def items_of(rows):
        return sorted([{
            "item": r.get("itemCd"), "name": r.get("itemNm"), "pgrp": r.get("prodGroup") or "-", "alois": r.get("aloisCd"),
            "qty": int(num(r.get("crtQty"))), "val": round(val_fn(r)),
        } for r in rows], key=lambda i: i["val"], reverse=True)

    oa = [r for r in pw_rows if r.get("aloisCd") == "OA"]
    ov = [r for r in pw_rows if r.get("aloisCd") == "OV"]
    oa_val, ov_val = round(sum(val_fn(r) for r in oa)), round(sum(val_fn(r) for r in ov))
    oa_qty = round(sum(num(r.get("crtQty")) for r in oa))
    ov_qty = round(sum(num(r.get("crtQty")) for r in ov))
    oa_items, ov_items = items_of(oa), items_of(ov)
    top_items = sorted(oa_items + ov_items, key=lambda i: i["val"], reverse=True)[:10]
    total = oa_val + ov_val
    return {
        "total": total, "inv_total": inv_total,
        "pct": round(100 * total / inv_total, 2) if inv_total else 0,
        "oa_val": oa_val, "ov_val": ov_val, "oa_lines": len(oa), "ov_lines": len(ov),
        "oa_qty": oa_qty, "ov_qty": ov_qty, "item_count": len(oa) + len(ov),
        "oa_items": oa_items, "ov_items": ov_items, "top_items": top_items,
    }


def build_longstock(pw_rows, inv_total, now):
    """lastPurcDt(최종입고일) 기준 12개월+/24개월+ 장기재고(입고 후 재구매 없음) 집계."""
    def val_fn(r):
        return num(r.get("crtQty")) * num(r.get("movPrc"))

    def months_since(dt_str):
        if not dt_str:
            return None
        d = datetime.strptime(dt_str[:10], "%Y-%m-%d")
        return (now.year - d.year) * 12 + (now.month - d.month) - (1 if now.day < d.day else 0)

    def items_of(rows):
        return sorted([{
            "item": r.get("itemCd"), "name": r.get("itemNm"), "pgrp": r.get("prodGroup") or "-", "alois": r.get("aloisCd"),
            "qty": int(num(r.get("crtQty"))), "val": round(val_fn(r)),
            "last_in": (r.get("lastPurcDt") or "")[:10] or "입고이력 없음",
        } for r in rows], key=lambda i: i["val"], reverse=True)

    m12_rows, m24_rows = [], []
    for r in pw_rows:
        months = months_since(r.get("lastPurcDt"))
        if months is None and val_fn(r) == 0:
            continue  # 입고이력 없고 금액도 0원인 품목은 리스트에서 제외
        if months is None or months >= 12:
            m12_rows.append(r)
        if months is None or months >= 24:
            m24_rows.append(r)

    m12_items, m24_items = items_of(m12_rows), items_of(m24_rows)
    m12_val = round(sum(i["val"] for i in m12_items))
    m24_val = round(sum(i["val"] for i in m24_items))
    return {
        "m12_items": m12_items, "m12_val": m12_val, "m12_count": len(m12_items),
        "m12_pct": round(100 * m12_val / inv_total, 2) if inv_total else 0,
        "m24_items": m24_items, "m24_val": m24_val, "m24_count": len(m24_items),
        "m24_pct": round(100 * m24_val / inv_total, 2) if inv_total else 0,
    }


def build_daily_stockcheck(recv_rows, pw_rows, today_str):
    """당일 입고된 부품 중 현재고>0인 것만 LOCATION 오름차순으로 나열 (일일 재고조사용 인쇄 리스트)."""
    today_items = {r.get("itemCd") for r in recv_rows if r.get("itemCd")}
    inv_by_item = {r.get("itemCd"): r for r in pw_rows}

    rows = []
    for item in today_items:
        r = inv_by_item.get(item)
        if not r:
            continue
        qty = int(num(r.get("crtQty")))
        if qty <= 0:
            continue  # 출고돼서 재고 0이 된 부품은 제외
        rows.append({
            "loc": r.get("lctCd") or "-", "item": item, "name": r.get("itemNm"),
            "alois": r.get("aloisCd") or "-", "qty": qty,
        })
    rows.sort(key=lambda x: (x["loc"] == "-", x["loc"], x["item"]))
    for i, r in enumerate(rows, 1):
        r["no"] = i

    return {"rows": rows, "count": len(rows), "date": today_str}


def build_weekly_stockcheck(turnover_rows, pw_rows, start_str, end_str):
    """지난주 월~토 출고(인보이스)된 부품 중 현재고>0인 것만 LOCATION 오름차순으로 나열 (이번주 월요일 주간 재고조사용)."""
    week_items = {r.get("itemCd") for r in turnover_rows if r.get("itemCd")}
    inv_by_item = {r.get("itemCd"): r for r in pw_rows}

    rows = []
    for item in week_items:
        r = inv_by_item.get(item)
        if not r:
            continue
        qty = int(num(r.get("crtQty")))
        if qty <= 0:
            continue  # 재고가 0인 부품은 제외
        rows.append({
            "loc": r.get("lctCd") or "-", "item": item, "name": r.get("itemNm"),
            "alois": r.get("aloisCd") or "-", "qty": qty,
        })
    rows.sort(key=lambda x: (x["loc"] == "-", x["loc"], x["item"]))
    for i, r in enumerate(rows, 1):
        r["no"] = i

    return {"rows": rows, "count": len(rows), "start": start_str, "end": end_str}


def build_o_parts(rows, now, pgrp_map=None):
    """ALOIS O계열 중 미처리 출고요청(RO/SB/SP)에 걸려있는 재고. 요청수량×이동평균단가 = 원가."""
    def val_fn(r):
        return num(r.get("reqQty")) * num(r.get("movPrc"))

    def dday_of(dt_str):
        if not dt_str:
            return "-"
        d = datetime.strptime(dt_str[:10], "%Y-%m-%d")
        days = (now.date() - d.date()).days
        return f"D+{days}일" if days >= 0 else f"D-{-days}일"  # 미래 예약분(SB 등)은 D-N일

    o_rows = [r for r in rows if str(r.get("aloisCd") or "").startswith("O")]
    items = sorted([{
        "item": r.get("partNo"), "name": r.get("itemNm"), "alois": r.get("aloisCd"),
        "pgrp": (pgrp_map or {}).get(r.get("partNo")) or "-",
        "ref_tp": r.get("refDocTp") or "기타", "ref_no": r.get("refDocNo") or "-",
        "req_qty": int(num(r.get("reqQty"))), "crt_qty": int(num(r.get("crtQty"))),
        "avail_qty": int(num(r.get("availQty"))), "req_dt": (r.get("reqDt") or "")[5:10],
        "req_dt_full": r.get("reqDt") or "", "req_brch": r.get("reqBrchNm") or "-", "val": round(val_fn(r)),
        "sa": r.get("saNm") or r.get("reqUsrNm") or "-", "dday": dday_of(r.get("reqDt")),
    } for r in o_rows], key=lambda i: i["req_dt_full"])  # 오래된 요청이 위로

    by_type = {}
    for i in items:
        by_type.setdefault(i["ref_tp"], []).append(i)
    group_order = ["RO", "SB", "SP"]
    groups = []
    for tp in group_order:
        if tp in by_type:
            groups.append({"code": tp, "rows": by_type.pop(tp)})
    for tp, rs in by_type.items():
        groups.append({"code": tp, "rows": rs})
    for g in groups:
        g["count"] = len(g["rows"])
        g["qty_sum"] = sum(i["req_qty"] for i in g["rows"])
        g["val"] = round(sum(i["val"] for i in g["rows"]))

    return {"items": items, "count": len(items), "total_val": round(sum(i["val"] for i in items)), "groups": groups}


def build_sb_parts(req_rows, resv_rows, now, farthest_first=False, date_field="resvDtime", exclude_zero_stock=False):
    """주어진 예약(SB) 목록 중 미처리 부품 요청이 걸려있는 부품. SB별 그룹, 예약일 빠른(오래된) 순. 원가=요청수량×이동평균단가.
    exclude_zero_stock=True면 현재고 0(아직 입고 안 되고 주문 중인 파트)인 라인은 제외."""
    resv = {r.get("resvNo"): r for r in resv_rows if r.get("resvNo")}
    by_sb = {}
    for r in req_rows:
        if r.get("refDocTp") != "SB" or r.get("refDocNo") not in resv:
            continue
        if exclude_zero_stock and num(r.get("crtQty")) <= 0:
            continue
        by_sb.setdefault(r["refDocNo"], []).append(r)

    def dday_of(dt_str):
        if not dt_str:
            return "-"
        d = datetime.strptime(dt_str[:10], "%Y-%m-%d")
        days = (now.date() - d.date()).days
        if days == 0:
            return "오늘"
        return f"D+{days}일" if days > 0 else f"D-{-days}일"

    groups = []
    for sb, rs in by_sb.items():
        info = resv[sb]
        resv_dt = info.get(date_field) or info.get("resvDtime") or ""
        parts = sorted([{
            "item": r.get("partNo"), "name": r.get("itemNm"), "alois": r.get("aloisCd"),
            "sa": info.get("chrgSaNm") or "-",
            "req_qty": int(num(r.get("reqQty"))), "crt_qty": int(num(r.get("crtQty"))),
            "avail_qty": int(num(r.get("availQty"))),
            "val": round(num(r.get("reqQty")) * num(r.get("movPrc"))),
        } for r in rs], key=lambda i: -i["val"])
        groups.append({
            "code": sb, "resv_full": resv_dt, "resv": resv_dt[5:10], "dday": dday_of(resv_dt),
            "sa": info.get("chrgSaNm") or "-", "rows": parts, "count": len(parts),
            "qty_sum": sum(p["req_qty"] for p in parts), "val": sum(p["val"] for p in parts),
        })
    groups.sort(key=lambda g: g["resv_full"], reverse=farthest_first)

    all_parts = [p for g in groups for p in g["rows"]]
    o_parts = [p for p in all_parts if str(p["alois"] or "").startswith("O")]
    return {
        "groups": groups, "sb_count": len(groups), "line_count": len(all_parts),
        "total_val": sum(p["val"] for p in all_parts),
        "o_lines": len(o_parts), "o_val": sum(p["val"] for p in o_parts),
        "resv_total": len(resv),
    }


def build_open_ro(rows, now):
    """인보이스 안 된 진행 RO 리스트(오래된 순). 금액 = 부품+공임 판매가 - 할인. 고객정보는 공개사이트라 제외."""
    def amt(r):
        return (num(r.get("itemSumSalePrice")) + num(r.get("lbrSaleAmt"))
                - num(r.get("itemDcAmt")) - num(r.get("lblDcPrice")))

    def dday_of(dt_str):
        if not dt_str:
            return "-"
        d = datetime.strptime(dt_str[:10], "%Y-%m-%d")
        return f"D+{(now.date() - d.date()).days}일"

    seen, items = set(), []
    for r in rows:
        ro = r.get("roNm")
        if not ro or ro in seen or r.get("roStatCd") in ("06", "07"):
            continue
        seen.add(ro)
        pub = r.get("publishDt") or r.get("csltStartDt") or ""
        items.append({
            "ro": ro, "stat": r.get("roStat") or "-", "stat_cd": r.get("roStatCd") or "",
            "svc": r.get("svcType") or "-", "sa": r.get("chrgSaNm") or "-",
            "pub": pub[:10], "pub_full": pub, "dday": dday_of(pub), "amt": round(amt(r)),
        })
    items.sort(key=lambda i: i["pub_full"])

    by_stat = {}
    for i in items:
        by_stat.setdefault(i["stat_cd"], {"stat": i["stat"], "count": 0, "amt": 0})
        by_stat[i["stat_cd"]]["count"] += 1
        by_stat[i["stat_cd"]]["amt"] += i["amt"]
    stat_list = [by_stat[c] for c in OPEN_RO_STAT_CODES if c in by_stat]

    return {
        "rows": items, "count": len(items), "total_amt": sum(i["amt"] for i in items),
        "by_stat": stat_list, "oldest": items[0]["dday"] if items else "-",
    }


def build_nonmng(pw_rows, inv_total):
    """현재고리스트(부품창고) 중 비관리부품(nonMngItemAtcYn=Y)으로 체크된 부품. 금액=현재고×이동평균단가(원가)."""
    def val_fn(r):
        return num(r.get("crtQty")) * num(r.get("movPrc"))

    rows = sorted([{
        "item": r.get("itemCd"), "name": r.get("itemNm"), "alois": r.get("aloisCd"),
        "qty": int(num(r.get("crtQty"))), "avail_qty": int(num(r.get("ableQty"))),
        "prc": round(num(r.get("movPrc"))), "val": round(val_fn(r)),
        "last_in": (r.get("lastPurcDt") or "")[:10] or "-",
        "last_sale": (r.get("lastSaleDt") or "")[:10] or "-",
    } for r in pw_rows if r.get("nonMngItemAtcYn") == "Y"], key=lambda i: -i["val"])
    # 원가 0원이면서 최종입고일 조회가 안 되는 부품은 제외
    rows = [i for i in rows if not (i["val"] == 0 and i["last_in"] == "-")]

    total_val = sum(i["val"] for i in rows)
    return {
        "rows": rows, "count": len(rows), "total_val": total_val,
        "qty_sum": sum(i["qty"] for i in rows),
        "pct": round(100 * total_val / inv_total, 2) if inv_total else 0,
    }


def build_no_location(pw_rows, inv_total):
    """현재고리스트(부품창고) 중 재고는 있는데 LOCATION이 공란이거나 N/STK로 찍힌 부품. 이동평균가 0원은 제외. 금액=현재고×이동평균단가(원가)."""
    def val_fn(r):
        return num(r.get("crtQty")) * num(r.get("movPrc"))

    def is_no_loc(r):
        loc = (r.get("lctCd") or "").strip()
        return loc == "" or loc.upper() == "N/STK"

    rows = sorted([{
        "item": r.get("itemCd"), "name": r.get("itemNm"), "alois": r.get("aloisCd"),
        "loc": (r.get("lctCd") or "").strip() or "(공란)",
        "qty": int(num(r.get("crtQty"))), "prc": round(num(r.get("movPrc"))), "val": round(val_fn(r)),
        "last_purc": (r.get("lastPurcDt") or "")[:10] or "-",
    } for r in pw_rows if is_no_loc(r) and num(r.get("crtQty")) > 0 and num(r.get("movPrc")) > 0], key=lambda i: -i["val"])

    total_val = sum(i["val"] for i in rows)
    return {
        "rows": rows, "count": len(rows), "total_val": total_val,
        "qty_sum": sum(i["qty"] for i in rows),
        "pct": round(100 * total_val / inv_total, 2) if inv_total else 0,
    }


def build_o_available(pw_rows):
    """ALOIS O계열 중 DMS 가용재고(ableQty) 기준 순수 가용분만.
    ableQty는 DMS가 RO/SB/SP 등 모든 홀드를 이미 반영해 계산한 값이라, 이걸 직접 써야
    "현재고 > 가용재고인데 전체수량이 가용으로 잡히는" 모순이 안 생김. 금액도 가용수량 기준."""
    def val_fn(r):
        return num(r.get("ableQty")) * num(r.get("movPrc"))

    o_rows = [
        r for r in pw_rows
        if str(r.get("aloisCd") or "").startswith("O") and num(r.get("ableQty")) > 0
    ]
    parts = sorted([{
        "item": r.get("itemCd"), "name": r.get("itemNm"), "pgrp": r.get("prodGroup") or "-", "alois": r.get("aloisCd"),
        "qty": int(num(r.get("crtQty"))), "avail_qty": int(num(r.get("ableQty"))),
        "val": round(val_fn(r)),
    } for r in o_rows], key=lambda i: i["val"], reverse=True)

    by_code = {}
    for p in parts:
        by_code.setdefault(p["alois"], []).append(p)
    groups = [{
        "code": code, "rows": rs, "count": len(rs),
        "qty_sum": sum(i["qty"] for i in rs), "avail_sum": sum(i["avail_qty"] for i in rs),
        "val": round(sum(i["val"] for i in rs)),
    } for code, rs in sorted(by_code.items())]

    return {"rows": parts, "groups": groups, "count": len(parts), "total_val": round(sum(i["val"] for i in parts))}


def build_o_daily_flow(recv_rows, turnover_rows, month_start, today_str):
    """이번달 O계열 파트 일자별 + 등급(ALOIS 코드)별 입고/출고 금액(원가 기준)."""
    in_by_day = defaultdict(float)
    in_by_code = defaultdict(float)
    for r in recv_rows:
        code = str(r.get("aloisCd") or "")
        if not code.startswith("O"):
            continue
        amt = num(r.get("purcAmt"))
        d = (r.get("realWhDt") or "")[:10]
        if d:
            in_by_day[d] += amt
        in_by_code[code] += amt

    out_by_day = defaultdict(float)
    out_by_code = defaultdict(float)
    for r in turnover_rows:
        code = str(r.get("aloisCd") or "")
        if not code.startswith("O"):
            continue
        amt = num(r.get("oriSumAmt"))
        d = (r.get("invDt") or "")[:10]
        if d:
            out_by_day[d] += amt
        out_by_code[code] += amt

    start_d = datetime.strptime(month_start, "%Y-%m-%d").date()
    end_d = datetime.strptime(today_str, "%Y-%m-%d").date()
    all_days = []
    d = start_d
    while d <= end_d:
        all_days.append(d.strftime("%Y-%m-%d"))
        d += timedelta(days=1)

    rows = []
    for d in all_days:
        in_val = round(in_by_day.get(d, 0.0))
        out_val = round(out_by_day.get(d, 0.0))
        rows.append({"date": d[5:10], "in_val": in_val, "out_val": out_val, "net": in_val - out_val})

    codes = sorted(set(in_by_code) | set(out_by_code))
    by_code = []
    for c in codes:
        in_val = round(in_by_code.get(c, 0.0))
        out_val = round(out_by_code.get(c, 0.0))
        by_code.append({"code": c, "in_val": in_val, "out_val": out_val, "net": in_val - out_val})

    total_in = round(sum(r["in_val"] for r in rows))
    total_out = round(sum(r["out_val"] for r in rows))
    in_days = sum(1 for r in rows if r["in_val"] > 0)
    avg_in = round(total_in / in_days) if in_days else 0
    return {
        "rows": rows, "total_in": total_in, "total_out": total_out, "net": total_in - total_out,
        "by_code": by_code, "in_days": in_days, "avg_in": avg_in,
    }


def _num_of(r):
    return r.get("parInvNo") or r.get("roNo") or ""


def build_ext_shop(rows, period_label):
    c_rows = [r for r in rows if r.get("calcTpCd") == "C"]

    def group_by(detl):
        grp = defaultdict(lambda: {"invTot": 0.0, "cost": 0.0, "margin": 0.0, "num": None, "cust": None, "dt": None})
        for r in c_rows:
            if r.get("calcDetlTpNm") != detl:
                continue
            g = grp[r.get("invNo")]
            g["invTot"] += num(r.get("invTotAmt"))
            g["cost"] += num(r.get("oriSumAmt"))
            g["margin"] += num(r.get("marginAmt"))
            g["num"] = g["num"] or _num_of(r)
            g["cust"] = g["cust"] or r.get("custNm")
            g["dt"] = g["dt"] or r.get("invDt")
        return grp

    ext_grp, shop_grp = group_by("외부"), group_by("외부공업사")

    def items_of(grp):
        return [{
            "inv": k, "num": g["num"], "cust": g["cust"], "dt": (g["dt"] or "")[:16],
            "amt": round(g["invTot"]), "cost": round(g["cost"]),
        } for k, g in grp.items()]

    ext_items, shop_items = items_of(ext_grp), items_of(shop_grp)
    ext_valid = [i for i in ext_items if i["amt"] != 0]
    ext_all = sorted(ext_valid, key=lambda i: i["dt"], reverse=True)
    ext_total = sum(i["amt"] for i in ext_items)
    ext_cost = sum(i["cost"] for i in ext_items)
    ext_margin = round(sum(num(r.get("marginAmt")) for r in c_rows if r.get("calcDetlTpNm") == "외부"))

    shop_sorted = sorted(shop_items, key=lambda i: i["dt"], reverse=True)
    for i in shop_sorted:
        i["cancelled"] = i["amt"] == 0
        i["status"] = "완료" if i["amt"] != 0 else "취소"
    shop_total = sum(i["amt"] for i in shop_items)
    shop_cost = sum(i["cost"] for i in shop_items)
    if shop_items and shop_total == 0:
        shop_note = f"발행 {len(shop_items)}건 · 전액 취소되어 순매출 없음"
    elif not shop_items:
        shop_note = "이번 달 발행 없음"
    else:
        shop_note = f"유효 {len([i for i in shop_items if i['amt'] != 0])}건"

    ext = {
        "period": period_label, "total": ext_total, "cost": ext_cost, "margin": ext_margin,
        "valid_count": len(ext_valid), "cancelled_count": len(ext_items) - len(ext_valid),
        "lines": len([r for r in c_rows if r.get("calcDetlTpNm") == "외부"]), "top5": ext_all,
    }
    shop = {
        "total": shop_total, "cost": shop_cost,
        "valid_count": len([i for i in shop_items if i["amt"] != 0]),
        "lines": len([r for r in c_rows if r.get("calcDetlTpNm") == "외부공업사"]),
        "invoices": shop_sorted, "foot_note": shop_note,
    }
    return ext, shop


def _prod_code(r):
    m = re.match(r"^\[(\d+)\]", r.get("prodType") or "")
    return m.group(1) if m else None


_STAFF_NAME_RE = re.compile(r"\((직원|영업직원)\)\s*$")
_BIZ_NAME_KEYWORDS = ("모터스", "상사", "공업사", "센터", "딜러", "산업", "상회", "무역", "수출", "(주)", "주식회사", "카센터", "정비")
_RO_LABOR_EXEMPT_KEYWORDS = (
    "오일", "워셔", "부동액", "요소수", "키", "배터리",  # 품목명이 한글인 경우
    "oil", "washer", "coolant", "antifreeze", "urea", "adblue", "key", "batter",  # 품목명이 영문인 경우 (DMS 기본 표기)
)
_EPC_EXCLUDE_ITEM_CODES = ("66126912985", "61319217643")  # 차량별 EPC 매칭이 필요 없는 범용 부품(키 배터리 등)


def build_ext_audit(rows, today_str):
    """외부판매(할인가) 오분류 의심 건 자동 적발 — '외부검증' 엑셀 매크로의 9단계 체크리스트를 그대로 룰로 옮긴 것.
    정산유형 C(고객) & 품목유형 부품/공임만 대상. 차대번호(VIN) 유무로 RO(정비연계)/SP(부품판매) 구분."""
    def is_staff_name(nm):
        return bool(_STAFF_NAME_RE.search((nm or "").strip()))

    def looks_like_business(nm):
        return any(k in (nm or "") for k in _BIZ_NAME_KEYWORDS)

    c_rows = [r for r in rows if r.get("calcTpCd") == "C" and r.get("itemTpCdNm") in ("부품", "공임")]

    # 인보이스번호 기준 합계금액이 0원(취소/반품으로 상쇄된 건)이면 애초에 분석 대상에서 제외
    inv_totals = defaultdict(float)
    for r in rows:
        inv_no = r.get("invNo")
        if inv_no:
            inv_totals[inv_no] += num(r.get("invTotAmt"))
    c_rows = [r for r in c_rows if round(inv_totals.get(r.get("invNo"), 1)) != 0]

    def base_of(r, reason):
        return {
            "inv": r.get("invNo") or "-", "ro": r.get("roNo") or "-", "spno": r.get("parInvNo") or "-",
            "cust": r.get("custNm") or "-",
            "vin": r.get("vinNo") or "-", "pgrp": _prod_code(r) or "-", "detl": r.get("calcDetlTpNm") or "-",
            "itemcd": r.get("itemCd") or "-", "item": r.get("itemNm") or "-", "amt": round(num(r.get("invTotAmt"))),
            "dt": (r.get("invDt") or "")[:10], "reason": reason,
        }

    # RO(정비연계)별 공임 존재 여부 — 같은 RO라도 인보이스가 나뉘어 공임은 보증(I) 등 다른 정산유형으로,
    # 부품만 고객(C)으로 청구되는 경우가 있어 정산유형 상관없이 전체 rows에서 공임 존재 여부를 확인해야 함
    ro_labor_total = defaultdict(float)
    for r in rows:
        if r.get("itemTpCdNm") == "공임" and r.get("roNo"):
            ro_labor_total[r.get("roNo")] += num(r.get("invTotAmt"))

    # RO(정비연계)별 부품 합계 (고객 청구분만) — 공임이 하나라도 있으면 그 RO는 "부품만 있는 건" 판정에서 제외
    by_ro = defaultdict(lambda: {"parts": 0.0, "part_names": [], "part_rows": [], "sample": None})
    for r in c_rows:
        ro = r.get("roNo")
        if not ro:
            continue
        g = by_ro[ro]
        amt = num(r.get("invTotAmt"))
        if r.get("itemTpCdNm") == "부품":
            g["parts"] += amt
            g["part_names"].append(r.get("itemNm") or "")
            g["part_rows"].append(r)
            g["sample"] = g["sample"] or r

    sp_violations, sp_review, ro_violations, ro_review, epc_check = [], [], [], [], []

    for r in c_rows:
        vin = (r.get("vinNo") or "").strip()
        detl = r.get("calcDetlTpNm") or ""
        cust = r.get("custNm") or ""
        pgrp = _prod_code(r)
        has_ro = bool(r.get("roNo"))

        if not has_ro:
            # ---- SP (부품판매) ----
            if pgrp != "7":
                if not vin and detl == "고객":
                    sp_violations.append(base_of(r, "차대 없는 일반부품인데 정산상세유형이 '고객' — 원래 외부판매 대상 (BMW 타 딜러 판매는 예외)"))
                elif not vin and detl in ("외부", "외부공업사") and not is_staff_name(cust) and not looks_like_business(cust):
                    sp_review.append(base_of(r, "차대 없는 외부/외부공업사 건인데 고객명이 직원/업체명 형식이 아님 — 확인 필요"))
            else:
                if not vin and detl == "고객" and not is_staff_name(cust):
                    sp_violations.append(base_of(r, "라이프스타일(부품그룹7) 차대없음 '고객' 건인데 고객명이 직원 형식이 아님"))
                elif not vin and detl in ("외부", "외부공업사") and not is_staff_name(cust):
                    sp_violations.append(base_of(r, "라이프스타일(부품그룹7) 차대없음 외부/외부공업사 건인데 고객명이 직원 형식이 아님 (마케팅 에이전시 예외 확인)"))
            if vin and detl in ("외부", "외부공업사"):
                sp_violations.append(base_of(r, "차대(VIN) 있는 외부/외부공업사 건 — 외부업체는 차대정보가 있으면 안 됨"))
            elif vin and detl == "고객" and looks_like_business(cust):
                sp_review.append(base_of(r, "차대 있는 '고객' 건인데 고객명이 업체명으로 추정 — 외부업체 여부 확인"))
            if vin and pgrp == "1" and detl == "고객" and r.get("itemTpCdNm") == "부품" and r.get("itemCd") not in _EPC_EXCLUDE_ITEM_CODES:
                epc_check.append(base_of(r, "EPC 확인"))
        else:
            # ---- RO (정비 연계) ----
            if detl in ("외부", "외부공업사"):
                ro_violations.append(base_of(r, "RO(정비연계) 건인데 정산상세유형이 외부/외부공업사 — RO는 '고객' 유형만 가능"))
            elif detl == "고객" and looks_like_business(cust) and round(ro_labor_total.get(r.get("roNo"), 0)) == 0:
                ro_review.append(base_of(r, "RO 건인데 공임 없이 부품만 존재 + 고객명이 업체명으로 추정 — 외부업체라면 SP로 처리했어야 함"))

    # RO 공임 0원 + 부품만 존재 (오일/워셔액 등 예외 제외)
    for ro, g in by_ro.items():
        if round(ro_labor_total.get(ro, 0)) == 0 and g["parts"] > 0 and g["sample"] is not None:
            if not all(any(k in nm.lower() for k in _RO_LABOR_EXEMPT_KEYWORDS) for nm in g["part_names"]):
                for pr in g["part_rows"]:
                    ro_review.append(base_of(pr, "공임 0원 + 부품만 존재 — 예외 품목(오일/워셔액/부동액/요소수/키/배터리) 아니면 확인 필요"))

    for lst in (sp_violations, sp_review, ro_violations, ro_review, epc_check):
        lst.sort(key=lambda i: i["dt"], reverse=True)

    return {
        "date": today_str,
        "sp_violations": sp_violations, "sp_review": sp_review,
        "ro_violations": ro_violations, "ro_review": ro_review,
        "epc_check": epc_check, "epc_count": len(epc_check),
        "epc_amt": round(sum(i["amt"] for i in epc_check)),
        "sp_v_count": len(sp_violations), "sp_r_count": len(sp_review),
        "ro_v_count": len(ro_violations), "ro_r_count": len(ro_review),
        "total_count": len(sp_violations) + len(sp_review) + len(ro_violations) + len(ro_review),
        "violation_count": len(sp_violations) + len(ro_violations),
        "review_count": len(sp_review) + len(ro_review),
        "violation_amt": round(sum(i["amt"] for i in sp_violations + ro_violations)),
        "review_amt": round(sum(i["amt"] for i in sp_review + ro_review)),
    }


def build_acc_tire(rows):
    ci_rows = [r for r in rows if r.get("calcTpCd") in ("C", "I")]
    acc_codes = {"3", "5", "7"}

    def group_by_invoice(codes):
        grp = defaultdict(lambda: {"invTot": 0.0, "cost": 0.0, "num": None, "sa": None, "dt": None})
        for r in ci_rows:
            if _prod_code(r) not in codes:
                continue
            k = r.get("invNo")
            if not k:
                continue
            g = grp[k]
            g["invTot"] += num(r.get("invTotAmt"))
            g["cost"] += num(r.get("oriSumAmt"))
            g["num"] = g["num"] or _num_of(r)
            g["sa"] = g["sa"] or r.get("emplNm")
            g["dt"] = g["dt"] or r.get("invDt")
        return grp

    acc_grp = group_by_invoice(acc_codes)
    tire_grp = group_by_invoice({"8"})

    def sorted_items(grp):
        return sorted(
            [{"inv": k, "num": g["num"], "sa": g["sa"] or "-", "dt": (g["dt"] or "")[:16],
              "amt": round(g["invTot"]), "cost": round(g["cost"])} for k, g in grp.items()],
            key=lambda i: i["dt"], reverse=True,
        )

    acc_by_type, acc_lines, acc_total = [], 0, 0.0
    for c in sorted(acc_codes):
        crows = [r for r in ci_rows if _prod_code(r) == c]
        v = sum(num(r.get("invTotAmt")) for r in crows)
        acc_by_type.append({"code": c, "label": PROD_LABELS.get(c, c), "lines": len(crows), "val": round(v)})
        acc_lines += len(crows)
        acc_total += v
    acc_cost = round(sum(num(r.get("oriSumAmt")) for r in ci_rows if _prod_code(r) in acc_codes))

    tire_rows = [r for r in ci_rows if _prod_code(r) == "8"]
    tire_total = round(sum(num(r.get("invTotAmt")) for r in tire_rows))
    tire_cost = round(sum(num(r.get("oriSumAmt")) for r in tire_rows))

    by_sa = defaultdict(lambda: {"qty": 0.0, "val": 0.0, "margin": 0.0})
    for r in tire_rows:
        sa = r.get("emplNm") or "(미지정)"
        by_sa[sa]["qty"] += num(r.get("qty"))
        by_sa[sa]["val"] += num(r.get("invTotAmt"))
        by_sa[sa]["margin"] += num(r.get("marginAmt"))
    sa_list = sorted(
        [{"sa": k, "qty": round(v["qty"]), "val": round(v["val"]), "margin": round(v["margin"])}
         for k, v in by_sa.items()],
        key=lambda x: x["val"], reverse=True,
    )

    acc = {
        "total": round(acc_total), "cost": acc_cost, "lines": acc_lines,
        "invoices_count": len(acc_grp), "by_type": acc_by_type, "recent5": sorted_items(acc_grp)[:5],
    }
    tire = {
        "total": tire_total, "cost": tire_cost, "lines": len(tire_rows),
        "invoices_count": len(tire_grp), "by_sa": sa_list,
        "qty_total": sum(x["qty"] for x in sa_list), "margin_total": sum(x["margin"] for x in sa_list),
        "all_invoices": sorted_items(tire_grp),
    }
    return acc, tire


# ============================================================
# 렌더 & 저장
# ============================================================

def _git(*args, check=True, timeout=None):
    return subprocess.run(
        ["git", *args], cwd=str(BASE_DIR), capture_output=True, text=True, encoding="utf-8",
        check=check, timeout=timeout,
    )


def publish_to_github(commit_message: str) -> bool:
    """docs/ 안의 최신 대시보드를 깃허브에 push. 저장소가 아직 없으면 초기화."""
    if not (BASE_DIR / ".git").exists():
        print("깃 저장소 초기화...")
        _git("init")
        _git("remote", "add", "origin", GIT_REMOTE_URL)
        gitignore = BASE_DIR / ".gitignore"
        if not gitignore.exists():
            gitignore.write_text("__pycache__/\noutput/\n*.pyc\n", encoding="utf-8")

    _git("add", "docs", ".gitignore", "partsbay.py", "template.html.j2", "tire_page.html.j2", "requirements.txt", check=False)
    status = _git("status", "--porcelain")
    if not status.stdout.strip():
        print("깃허브에 올릴 변경사항이 없습니다 (이전과 동일한 데이터).")
        return True

    _git("commit", "-m", commit_message)
    print("깃허브에 push 중... (세션 만료 시 브라우저 로그인 창이 뜰 수 있습니다)")
    push_timeout = 60  # 인증 대기 등으로 상시 실행 루프가 무한정 멈추지 않게
    try:
        result = _git("push", "-u", "origin", "main", check=False, timeout=push_timeout)
        if result.returncode != 0:
            # main이 처음이라 브랜치 이름이 다를 수 있음 (master 등) -> 강제로 main 사용
            _git("branch", "-M", "main", check=False)
            result = _git("push", "-u", "origin", "main", check=False, timeout=push_timeout)
    except subprocess.TimeoutExpired:
        print("push가 시간 내에 끝나지 않았습니다 (인증 대기 중일 수 있음). 다음 주기에 재시도합니다.")
        return False
    if result.returncode != 0:
        print("push 실패:")
        print(result.stdout)
        print(result.stderr)
        return False
    print("깃허브 push 완료.")
    return True


def render(data: dict) -> str:
    env = Environment(loader=FileSystemLoader(str(BASE_DIR)), autoescape=False)
    template = env.get_template(TEMPLATE_NAME)
    data = dict(data)
    data.setdefault("tires", {})
    data["tires_json"] = json.dumps(data["tires"], ensure_ascii=False)
    return template.render(**data)


TIRE_TEMPLATE_NAME = "tire_page.html.j2"


def render_tire_page(tires: dict, generated_at: str) -> str:
    """타이어 사이즈 조회만 담은 독립 페이지 (다른 대시보드 데이터는 포함하지 않음)."""
    env = Environment(loader=FileSystemLoader(str(BASE_DIR)), autoescape=False)
    template = env.get_template(TIRE_TEMPLATE_NAME)
    return template.render(tires_json=json.dumps(tires or {}, ensure_ascii=False), generated_at=generated_at)


LOCK_FILE = BASE_DIR / ".partsbay.lock"
CYCLE_INTERVAL_SEC = 600  # 상시 실행 모드에서 갱신 주기 (10분)
ONCE = "--once" in sys.argv  # 테스트용: 한 번만 돌고 종료 (기본은 계속 켜져 있는 상시 실행)


def _pid_alive(pid: int) -> bool:
    """Windows에선 os.kill(pid, 0)이 생존여부와 무관하게 항상 예외를 던져서 못 씀
    (그래서 잠금이 사실상 계속 무시되고 있었음) -> tasklist로 직접 확인."""
    try:
        result = subprocess.run(
            ["tasklist", "/FI", f"PID eq {pid}", "/NH"],
            capture_output=True, text=True, timeout=5,
        )
        return str(pid) in result.stdout
    except Exception:
        return False  # 확인 불가하면 안전하게 "죽은 것"으로 간주해 새로 진행


def _acquire_lock() -> bool:
    """이미 실행 중인 인스턴스가 있으면 False. (중복 실행 -> 동시 로그인으로 세션 끊기는 사고 방지)"""
    if LOCK_FILE.exists():
        try:
            pid = int(LOCK_FILE.read_text().strip())
        except ValueError:
            pid = None
        if pid is not None and _pid_alive(pid):
            return False  # 여전히 실행 중
    LOCK_FILE.write_text(str(os.getpid()))
    return True


def load_reasons() -> dict:
    with REASONS_LOCK:
        try:
            return json.loads(REASONS_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}


def _open_reasons(data: dict) -> dict:
    """현재 진행RO 목록에 있는 RO의 사유만 사이트에 실음."""
    ros = {r["ro"] for r in data.get("openro", {}).get("rows", [])}
    return {k: v for k, v in load_reasons().items() if k in ros}


def update_reasons(changes: dict) -> dict:
    """{RO번호: 사유} 반영(빈 문자열이면 삭제). 잘못된 항목은 무시."""
    with REASONS_LOCK:
        try:
            cur = json.loads(REASONS_FILE.read_text(encoding="utf-8"))
        except Exception:
            cur = {}
        for ro, text in changes.items():
            if not isinstance(ro, str) or not re.fullmatch(r"RO\d{6,14}", ro) or not isinstance(text, str):
                continue
            text = text.strip()[:300]
            if text:
                cur[ro] = text
            else:
                cur.pop(ro, None)
        REASONS_FILE.write_text(json.dumps(cur, ensure_ascii=False, indent=2), encoding="utf-8")
        return cur


def republish_reasons() -> None:
    """사유만 바뀐 경우: 마지막 데이터로 다시 렌더해서 docs에 쓰고 push."""
    with PUBLISH_LOCK:
        if not LAST_DATA:
            return
        data = dict(LAST_DATA)
        data["reasons"] = _open_reasons(data)
        data["tires"] = load_tires()
        data["dday"] = build_dday(load_dday_items(), datetime.now(ZoneInfo("Asia/Seoul")))
        html = render(data)
        DOCS_DIR.mkdir(parents=True, exist_ok=True)
        (DOCS_DIR / "index.html").write_text(html, encoding="utf-8")
        (DOCS_DIR / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        (DOCS_DIR / "tire.html").write_text(render_tire_page(data["tires"], data.get("meta", {}).get("generated_at", "")), encoding="utf-8")
        if PUBLISH_TO_GITHUB:
            publish_to_github("진행RO 사유 갱신")


def _norm_tire_size(size) -> str:
    """'225/40R18 92Y' -> '2254018' (엑셀 필터링 열과 동일한 정규화: 속도/하중지수 제거, ZR/R/슬래시 제거)"""
    s = (str(size) if size else "").split(" ")[0]
    s = s.replace("ZR", "").replace("R", "").replace("/", "")
    return s


def load_tires() -> dict:
    """엑셀을 못 읽었을 때 쓸 마지막 성공본(캐시)."""
    with TIRE_LOCK:
        try:
            return json.loads(TIRE_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}


def import_tire_excel() -> dict | None:
    """타이어_마스터파일.xlsx의 '타이어 리스트' 시트를 읽어 tire_master.json 형식으로 변환.
    실패(파일 없음/엑셀에서 편집 중이라 잠김 등)하면 None을 반환해 이전 값을 그대로 쓰게 한다."""
    if not TIRE_XLSX_PATH.exists():
        return None
    try:
        wb = openpyxl.load_workbook(TIRE_XLSX_PATH, data_only=True, read_only=True)
        ws = wb["타이어 리스트"] if "타이어 리스트" in wb.sheetnames else wb.worksheets[0]
    except Exception as e:
        print(f"[타이어 엑셀 읽기 실패] {e}")
        return None

    tires = {}
    try:
        for row in ws.iter_rows(min_row=2):
            pn = row[1].value  # B열: 부품번호
            if not pn:
                continue
            pn = str(pn).strip()
            brand = row[3].value
            sale_tp = row[5].value  # F열: 판매유형
            # Yokohama는 D+1만 사용(D+0 제외), 나머지 브랜드는 D+0만 사용(D+1 제외)
            if brand == "Yokohama":
                if sale_tp != "D+1":
                    continue
            elif sale_tp == "D+1":
                continue
            size = row[6].value  # G열: 사이즈
            price = row[11].value  # L열: 구매단가
            try:
                price = round(float(price)) if price not in (None, "") else None
            except Exception:
                price = None
            tires[pn] = {
                "name": row[2].value, "brand": brand,
                "mat_no": row[4].value, "sale_tp": sale_tp, "size": size,
                "std_rft": row[7].value, "sw": row[8].value,
                "oe": row[9].value, "qr": row[10].value, "price": price,
                "bridge": (row[12].value or "") if len(row) > 12 else "", "filter_key": _norm_tire_size(size),
            }
    except Exception as e:
        print(f"[타이어 엑셀 파싱 실패] {e}")
        return None

    with TIRE_LOCK:
        TIRE_FILE.write_text(json.dumps(tires, ensure_ascii=False, indent=2), encoding="utf-8")
    return tires


def load_dday_items() -> list:
    """구글 시트를 못 읽었을 때 쓸 마지막 성공본(캐시)."""
    with DDAY_LOCK:
        try:
            return json.loads(DDAY_FILE.read_text(encoding="utf-8"))
        except Exception:
            return []


def import_dday_sheet() -> list | None:
    """구글 폼 응답 시트(타임스탬프/제목/목표일/메모)를 CSV로 읽어와 디데이 항목으로 변환.
    실패(연동 안 함/네트워크 오류 등)하면 None을 반환해 이전 값을 그대로 쓰게 한다."""
    if not DDAY_SHEET_CSV_URL:
        return None
    try:
        req = urllib.request.Request(DDAY_SHEET_CSV_URL, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            raw = resp.read().decode("utf-8-sig")
    except Exception as e:
        print(f"[디데이 시트 읽기 실패] {e}")
        return None

    items = []
    try:
        reader = csv.reader(io.StringIO(raw))
        next(reader, None)  # 헤더(타임스탬프/제목/목표일/메모) 건너뜀
        for r in reader:
            if len(r) < 3:
                continue
            title = (r[1] or "").strip()
            date_raw = (r[2] or "").strip()
            memo = (r[3] if len(r) > 3 else "").strip()
            if not title or not date_raw:
                continue
            nums = re.findall(r"\d+", date_raw)
            if len(nums) < 3:
                continue
            try:
                target = datetime(int(nums[0]), int(nums[1]), int(nums[2]))
            except Exception:
                continue
            items.append({"title": title, "target": target.strftime("%Y-%m-%d"), "memo": memo})
    except Exception as e:
        print(f"[디데이 시트 파싱 실패] {e}")
        return None

    with DDAY_LOCK:
        DDAY_FILE.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    return items


def build_dday(items: list, today: datetime) -> dict:
    """목표일까지 D-N(미래)/D-DAY(당일) 계산, 지난 디데이(D+1 이후)는 제외하고 임박 순으로 정렬."""
    rows = []
    for it in items:
        try:
            target = datetime.strptime(it["target"], "%Y-%m-%d")
        except Exception:
            continue
        delta = (target.date() - today.date()).days
        if delta < 0:
            continue
        dday = f"D-{delta}" if delta > 0 else "D-DAY"
        rows.append({"title": it.get("title", ""), "target": it["target"], "memo": it.get("memo", ""),
                      "dday": dday, "sort_key": delta})
    rows.sort(key=lambda x: x["sort_key"])
    return {"rows": rows, "count": len(rows), "linked": bool(DDAY_SHEET_CSV_URL)}


class ReasonHandler(BaseHTTPRequestHandler):
    def _send(self, code: int, body: dict | None = None) -> None:
        payload = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else b""
        self.send_response(code)
        origin = self.headers.get("Origin", "")
        if origin in ALLOWED_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Private-Network", "true")
        if body is not None:
            self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def do_OPTIONS(self):
        self._send(204)

    def _send_file(self, path: Path) -> None:
        data = path.read_bytes()
        ctype = mimetypes.guess_type(str(path))[0] or "application/octet-stream"
        if ctype.startswith("text/"):
            ctype += "; charset=utf-8"
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        route = self.path.split("?", 1)[0]
        if route.startswith("/ping"):
            self._send(200, {"ok": True})
        elif route == "/stock":
            origin = self.headers.get("Origin", "")
            if origin and origin not in ALLOWED_ORIGINS:
                self._send(403, {"ok": False, "error": "forbidden origin"})
                return
            qs = urllib.parse.parse_qs(self.path.split("?", 1)[1]) if "?" in self.path else {}
            pn = (qs.get("pn", [""])[0] or "").strip().upper()
            if not re.fullmatch(r"[A-Z0-9]{5,20}", pn):
                self._send(400, {"ok": False, "error": "pn은 영문/숫자 5~20자리 부품번호여야 합니다 (예: /stock?pn=80165B70CF2)"})
                return
            holder, done = {}, threading.Event()
            STOCK_QUEUE.put((pn, holder, done))
            if not done.wait(240):
                self._send(503, {"ok": False, "error": "데이터 갱신 중이라 시간 초과 - 잠시 후 다시 시도하세요"})
            elif "error" in holder:
                self._send(500, {"ok": False, "error": holder["error"]})
            else:
                self._send(200, {"ok": True, **holder["result"]})
        elif route.startswith("/reasons"):
            self._send(200, load_reasons())
        else:
            # 이 PC 전용 로컬 대시보드: docs 폴더(=깃허브에 올라가는 것과 동일)를 그대로 서빙 + 사유 편집 활성화
            rel = "index.html" if route in ("/", "") else route.lstrip("/")
            target = (DOCS_DIR / rel).resolve()
            if DOCS_DIR.resolve() in target.parents and target.is_file():
                self._send_file(target)
            else:
                self._send(404, {"ok": False})

    def do_POST(self):
        if not self.path.startswith("/reasons"):
            self._send(404, {"ok": False})
            return
        try:
            length = min(int(self.headers.get("Content-Length", "0")), 200_000)
            body = json.loads(self.rfile.read(length).decode("utf-8"))
            changes = body.get("set", {})
            if not isinstance(changes, dict):
                raise ValueError("set must be object")
        except Exception:
            self._send(400, {"ok": False})
            return
        cur = update_reasons(changes)
        threading.Thread(target=republish_reasons, daemon=True).start()
        self._send(200, {"ok": True, "reasons": cur})

    def log_message(self, *args):
        pass


def start_reason_server() -> None:
    try:
        server = ThreadingHTTPServer(("127.0.0.1", REASON_PORT), ReasonHandler)
    except OSError as e:
        print(f"[사유 입력 서버 시작 실패] 포트 {REASON_PORT}: {e}")
        return
    threading.Thread(target=server.serve_forever, daemon=True).start()
    print(f"진행RO 사유 / 타이어 마스터 입력 서버 시작 - 이 PC에서 http://127.0.0.1:{REASON_PORT} 로 접속하면 입력/수정 가능")


def run_cycle(page: Page, rr_page: Page | None = None) -> None:
    """이미 로그인된 page로 데이터 한 번 뽑아서 대시보드 생성 + 깃허브 push."""
    now = datetime.now(ZoneInfo("Asia/Seoul"))
    today_str = now.strftime("%Y-%m-%d")
    month_start = now.strftime("%Y-%m-01")
    period_label = f"{month_start} ~ {today_str}"
    generated_at = f"{now.month}월 {now.day}일 ({WEEKDAYS_KO[now.weekday()]}) · {now.strftime('%H:%M')} 기준"
    this_monday = now.date() - timedelta(days=now.weekday())
    last_monday = this_monday - timedelta(days=7)
    last_saturday = last_monday + timedelta(days=5)
    last_week_start_str = last_monday.strftime("%Y-%m-%d")
    last_week_end_str = last_saturday.strftime("%Y-%m-%d")

    print(" - 오늘 입고현황")
    recv_rows = extract_receiving(page, today_str)
    print(" - 현재고리스트")
    inv_rows = extract_inventory(page)
    print(" - Turn Over 리포트 (당월)")
    to_rows = extract_turnover(page, month_start, today_str)
    print(" - 출고요청관리 (O계열 RO/SB/SP)")
    req_rows = extract_part_requests(page)
    print(" - RO 리포트 (진행 RO: 인보이스 미완료)")
    open_ro_rows = extract_open_ros(page, today_str)
    print(" - 서비스예약현황 (노쇼 SB)")
    noshow_rows = extract_resv_sbs(page, "04", "2020-01-01")
    print(" - 서비스예약현황 (예약접수 SB, 오늘~)")
    sbresv_rows = extract_resv_sbs(page, "01", today_str)
    print(" - 차량접수 SB 중 RO 미발행 + 부품 (서비스예약현황 + RO 리포트)")
    sbcar_rows, sbcar_total = extract_carin_sbs_without_ro(page, req_rows, datetime.now(ZoneInfo("Asia/Seoul")).date())
    print(" - 입고현황 (당월, O파트 입출고 내역용)")
    recv_month_rows = extract_receiving_range(page, month_start, today_str)
    print(" - Turn Over 리포트 (지난주 월~토, 주간 재고조사용)")
    to_lastweek_rows = extract_turnover(page, last_week_start_str, last_week_end_str)

    print("집계 중...")
    recv = build_recv(recv_rows, today_str)
    inv, pw_rows = build_inventory(inv_rows)
    oaov = build_oaov(pw_rows, inv["total"])
    ext, shop = build_ext_shop(to_rows, period_label)
    extaudit = build_ext_audit(to_rows, today_str)
    acc, tire = build_acc_tire(to_rows)
    longstock = build_longstock(pw_rows, inv["total"], now)
    stockcheck = build_daily_stockcheck(recv_rows, pw_rows, today_str)
    stockcheck_week = build_weekly_stockcheck(to_lastweek_rows, pw_rows, last_week_start_str, last_week_end_str)
    pgrp_map = {r.get("itemCd"): r.get("prodGroup") for r in inv_rows if r.get("prodGroup")}
    opart = build_o_parts(req_rows, now, pgrp_map)
    oavail = build_o_available(pw_rows)
    nonmng = build_nonmng(pw_rows, inv["total"])
    noloc = build_no_location(pw_rows, inv["total"])
    openro = build_open_ro(open_ro_rows, now)
    noshow = build_sb_parts(req_rows, noshow_rows, now, exclude_zero_stock=True)
    sbresv = build_sb_parts(req_rows, sbresv_rows, now, farthest_first=True)
    sbcar = build_sb_parts(req_rows, sbcar_rows, now, date_field="carAcptDtime")
    sbcar["resv_total"] = sbcar_total
    oflow = build_o_daily_flow(recv_month_rows, to_rows, month_start, today_str)
    calendar_image = next((f for f in CALENDAR_IMAGE_CANDIDATES if (DOCS_DIR / f).exists()), None)

    data = {
        "meta": {"branch_name": BRANCH_NAME, "brch_code": f"BRCH {BRCH_CD}", "generated_at": generated_at,
                 "calendar_image": calendar_image},
        "recv": recv, "inv": inv, "oaov": oaov, "ext": ext, "shop": shop, "acc": acc, "tire": tire,
        "longstock": longstock, "opart": opart, "oavail": oavail, "oflow": oflow, "openro": openro, "noshow": noshow, "sbresv": sbresv, "sbcar": sbcar, "nonmng": nonmng,
        "stockcheck": stockcheck, "stockcheck_week": stockcheck_week, "noloc": noloc, "extaudit": extaudit,
    }

    global LAST_DATA
    with PUBLISH_LOCK:
        LAST_DATA = dict(data)
        data["reasons"] = _open_reasons(data)
        data["tires"] = import_tire_excel()
        if data["tires"] is None:
            data["tires"] = load_tires()
        dday_items = import_dday_sheet()
        if dday_items is None:
            dday_items = load_dday_items()
        data["dday"] = build_dday(dday_items, now)
        html = render(data)
        stamp = now.strftime("%Y%m%d_%H%M")
        out_path = OUTPUT_DIR / f"dashboard_{stamp}.html"
        out_path.write_text(html, encoding="utf-8")
        (OUTPUT_DIR / "latest.html").write_text(html, encoding="utf-8")
        data_json = json.dumps(data, ensure_ascii=False, indent=2)
        (OUTPUT_DIR / "latest.json").write_text(data_json, encoding="utf-8")

        DOCS_DIR.mkdir(parents=True, exist_ok=True)
        (DOCS_DIR / "index.html").write_text(html, encoding="utf-8")
        (DOCS_DIR / "data.json").write_text(data_json, encoding="utf-8")
        (DOCS_DIR / "tire.html").write_text(render_tire_page(data["tires"], generated_at), encoding="utf-8")

        print(f"완료: {out_path}")
        if ONCE:
            webbrowser.open(out_path.resolve().as_uri())

        if PUBLISH_TO_GITHUB:
            publish_to_github(f"대시보드 갱신 {generated_at}")

    # 매출 대시보드(지점장 전용, 비밀번호 암호화 페이지) — DMS 세션을 두 개 띄울 수 없어서 별도 데몬이 아니라 여기서 이어서 실행.
    # 실패해도 파츠베이 갱신에는 영향 없음.
    if SALES_ENABLED:
        print(" - 매출 대시보드")
        try:
            if str(SALES_DIR) not in sys.path:
                sys.path.insert(0, str(SALES_DIR))
            import importlib
            import sales
            importlib.reload(sales)  # sales.py 를 고쳐도 데몬 재시작(=DMS 재로그인) 없이 다음 주기부터 반영
            sales.run(to_rows, lambda s, e: extract_turnover(page, s, e), now, BRANCH_NAME, page=page, rr_page=rr_page)
        except Exception as e:
            print(f"   [매출 대시보드 실패, 다음 주기에 재시도] {type(e).__name__}: {e}")


def main():
    if not _acquire_lock():
        print("이미 실행 중인 인스턴스가 있습니다. (같은 프로그램을 두 번 띄우면 DMS 세션이 끊길 수 있어 종료합니다)")
        return
    AUTH_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    start_reason_server()
    try:
        print("동성모터스 PARTS — DMS 접속 중...")
        with sync_playwright() as p:
            # 브라우저는 한 번만 띄워서 계속 켜둔다(로그인 세션 유지). 화면엔 안 보이게
            # 항상 최소화 상태로 띄우고, 로그인이 필요할 때만 notify_user()가 알려준다.
            ctx = p.chromium.launch_persistent_context(
                str(AUTH_DIR), headless=False, viewport={"width": 1440, "height": 900},
                args=["--start-minimized"],
            )
            page = ctx.new_page()
            ensure_logged_in(page)

            if ONCE:
                print("로그인 확인 완료. 데이터를 추출합니다...")
                run_cycle(page)
                ctx.close()
                return

            # 롤스로이스 DMS(매출 대시보드 RR 항목용) — 도메인이 달라 BMW 세션과 충돌하지 않음. 로그인은 그 창에서 직접,
            # 로그인 전이어도 BMW 갱신은 그대로 진행되고 매출 모듈이 매 주기 로그인 여부만 확인한다.
            rr_page = None
            if SALES_ENABLED and RR_ENABLED:
                try:
                    RR_AUTH_DIR.mkdir(parents=True, exist_ok=True)
                    rr_ctx = p.chromium.launch_persistent_context(
                        str(RR_AUTH_DIR), headless=False, viewport={"width": 1440, "height": 900},
                        args=["--start-minimized"],
                    )
                    rr_page = rr_ctx.new_page()
                    rr_page.goto(f"{RR_BASE_URL}/selectHome.do")
                except Exception as e:
                    print(f"[RR DMS 브라우저 시작 실패 — 매출 RR 항목만 빠짐] {type(e).__name__}: {e}")
                    rr_page = None

            print(f"로그인 확인 완료. 상시 실행 시작 — {CYCLE_INTERVAL_SEC//60}분마다 자동 갱신합니다 (창은 계속 켜둔 채 백그라운드로 동작).")
            while True:
                try:
                    ensure_logged_in(page)  # 그 사이 세션이 끊겼으면 재확인
                    run_cycle(page, rr_page)
                except Exception as e:
                    print(f"[이번 주기 실패, 다음 주기에 재시도] {type(e).__name__}: {e}")
                print(f"다음 갱신까지 {CYCLE_INTERVAL_SEC}초 대기...")
                deadline = time.time() + CYCLE_INTERVAL_SEC
                while time.time() < deadline:  # 대기 중에도 /stock 부품 재고조회 요청을 처리
                    serve_stock_requests(page)
                    try:
                        page.wait_for_timeout(1000)
                    except Exception:
                        time.sleep(1)
    finally:
        LOCK_FILE.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
